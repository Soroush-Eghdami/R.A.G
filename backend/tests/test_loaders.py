"""Loader tests."""

from __future__ import annotations

import pymupdf
import pytest
from docx import Document
from rag.loaders import EmptyDocument, LoaderError, UnsupportedFileType, load_document


def test_txt_reads_utf8(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello world", encoding="utf-8")
    blocks = load_document(path)
    assert len(blocks) == 1
    assert blocks[0].text == "hello world"
    assert blocks[0].page is None


def test_txt_falls_back_to_cp1252(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes("caf\xe9".encode("latin-1"))
    assert load_document(path)[0].text == "café"


def test_markdown_keeps_heading_trail(tmp_path):
    path = tmp_path / "guide.md"
    path.write_text(
        "# Setup\n\nInstall it.\n\n## Docker\n\nRun the image.\n",
        encoding="utf-8",
    )
    blocks = load_document(path)
    assert [block.heading for block in blocks] == ["Setup", "Setup > Docker"]
    assert blocks[1].text == "Run the image."


def test_markdown_strips_inline_noise(tmp_path):
    path = tmp_path / "readme.md"
    path.write_text("Use **bold** and [`code`](http://x).\n", encoding="utf-8")
    assert load_document(path)[0].text == "Use bold and code."


def test_pdf_keeps_page_numbers(tmp_path):
    path = tmp_path / "doc.pdf"
    document = pymupdf.open()
    for number in (1, 2):
        page = document.new_page()
        page.insert_text((72, 72), f"page {number} content")
    document.save(path)
    document.close()

    blocks = load_document(path)
    assert [block.page for block in blocks] == [1, 2]
    assert "page 2 content" in blocks[1].text


def test_docx_tracks_headings_and_tables(tmp_path):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Overview", level=1)
    document.add_paragraph("Body text.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "key"
    table.rows[0].cells[1].text = "value"
    document.save(path)

    blocks = load_document(path)
    assert blocks[0].text == "Body text."
    assert blocks[0].heading == "Overview"
    assert "key | value" in blocks[1].text


def test_unsupported_suffix(tmp_path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04")
    with pytest.raises(UnsupportedFileType):
        load_document(path)


def test_missing_file():
    with pytest.raises(LoaderError):
        load_document("does-not-exist.txt")


def test_empty_pdf_is_reported(tmp_path):
    path = tmp_path / "blank.pdf"
    document = pymupdf.open()
    document.new_page()
    document.save(path)
    document.close()
    with pytest.raises(EmptyDocument):
        load_document(path)
